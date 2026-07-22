from __future__ import annotations

import csv
import json
import re
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable

import openpyxl


BASE_DIR = Path(__file__).resolve().parent
NS = {"sii": "http://www.sii.cl/SiiDte"}


def text_key(value: str | None) -> str:
    if value is None:
        return ""
    return re.sub(r"[^a-z0-9]+", "", str(value).lower())


def find_key(keys: Iterable[str], *fragments: str) -> str:
    normalized = {key: text_key(key) for key in keys}
    wanted = [text_key(fragment) for fragment in fragments]
    for key, norm in normalized.items():
        if all(fragment in norm for fragment in wanted):
            return key
    raise KeyError(f"No key for fragments {fragments}. Keys: {list(keys)}")


def parse_number(value: object) -> float:
    if value in (None, "", "nan"):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    raw = str(value).strip().replace(".", "").replace(",", ".")
    try:
        return float(raw)
    except ValueError:
        return 0.0


def parse_date(value: object) -> str | None:
    if value in (None, ""):
        return None
    raw = str(value).strip()
    for fmt in (
        "%d/%m/%Y %H:%M:%S",
        "%d/%m/%Y",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d",
        "%d-%m-%Y",
    ):
        try:
            return datetime.strptime(raw, fmt).date().isoformat()
        except ValueError:
            pass
    return raw[:10]


def read_csv_rows(pattern: str, delimiter: str) -> Iterable[tuple[Path, dict[str, str]]]:
    for path in sorted(BASE_DIR.glob(pattern)):
        with path.open("r", encoding="latin1", errors="replace", newline="") as handle:
            reader = csv.DictReader(handle, delimiter=delimiter)
            for row in reader:
                yield path, row


@dataclass
class XmlDoc:
    file_name: str
    dte_type: str | None
    folio: str | None
    issue_date: str | None
    total: float
    net: float
    vat: float
    commission_net: float
    commission_vat: float
    recipient_name: str | None
    detail_count: int


def load_th() -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for path, row in read_csv_rows("TH/*.csv", ";"):
        order_key = find_key(row.keys(), "pedido")
        invoice_key = find_key(row.keys(), "factura")
        cycle_key = find_key(row.keys(), "ciclo", "factur")
        type_key = find_key(row.keys(), "tipo")
        amount_key = find_key(row.keys(), "importe")
        item_key = find_key(row.keys(), "posici", "pedido")
        sku_key = find_key(row.keys(), "sku", "oferta")
        created_key = find_key(row.keys(), "creaci")
        rows.append(
            {
                "source_file": path.name,
                "order_id": (row.get(order_key) or "").strip(),
                "invoice_number": (row.get(invoice_key) or "").strip(),
                "cycle_date": parse_date(row.get(cycle_key)),
                "transaction_type": (row.get(type_key) or "").strip(),
                "amount": parse_number(row.get(amount_key)),
                "order_line_id": (row.get(item_key) or "").strip(),
                "offer_sku": (row.get(sku_key) or "").strip(),
                "created_date": parse_date(row.get(created_key)),
            }
        )
    return rows


def load_ciclos() -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for path, row in read_csv_rows("CICLOS/*.csv", ";"):
        invoice_key = find_key(row.keys(), "factura")
        amount_key = find_key(row.keys(), "amount", "transferred")
        rows.append(
            {
                "source_file": path.name,
                "order_id": (row.get("Order number") or "").strip(),
                "invoice_number": (row.get(invoice_key) or "").strip(),
                "order_status": (row.get("Estado del pedido") or "").strip(),
                "offer_sku": (row.get("Offer SKU") or "").strip(),
                "transfer_amount": parse_number(row.get(amount_key)),
                "created_date": parse_date(row.get("Date created")),
            }
        )
    return rows


def load_ff() -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for path, row in read_csv_rows("FF/[0-9]*_ff_2815.csv", ","):
        rows.append(
            {
                "source_file": path.name,
                "order_id": (row.get("order_id") or "").strip(),
                "order_line_id": (row.get("order_line_id") or "").strip(),
                "order_category_type": (row.get("order_category_type") or "").strip(),
                "order_amount": parse_number(row.get("order_amount")),
                "commission_fee": parse_number(row.get("commission_fee")),
                "transfer_amount": parse_number(row.get("transfer_amount")),
                "accounting_document_date": parse_date(row.get("accounting_document_creation_date")),
            }
        )
    return rows


def load_ff_adjustments() -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for path, row in read_csv_rows("FF/abonos_descuentos_*.csv", ","):
        adjustment_total = 0.0
        for key, value in row.items():
            if key in {
                "accounting_document_creation_date",
                "currency_iso_code",
                "date_created",
                "order_id",
                "order_category_type",
                "shipping_type_code",
                "shop_id",
                "shop_name",
                "refund_commission_fee",
                "refund_order_amount",
                "commission_fee",
                "order_amount",
                "transfer_amount",
            }:
                continue
            adjustment_total += parse_number(value)
        rows.append(
            {
                "source_file": path.name,
                "order_id": (row.get("order_id") or "").strip(),
                "adjustment_total": adjustment_total,
                "accounting_document_date": parse_date(row.get("accounting_document_creation_date")),
            }
        )
    return rows


def load_seller() -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for path in sorted(BASE_DIR.glob("SELLER/*.xlsx")):
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        iterator = sheet.iter_rows(values_only=True)
        headers = list(next(iterator))
        liquidation_idx = next(i for i, value in enumerate(headers) if value and "liquid" in str(value).lower())
        order_idx = next(i for i, value in enumerate(headers) if value and "orden de compra" in str(value).lower())
        payable_idx = next(i for i, value in enumerate(headers) if value == "A pagar")
        for row in iterator:
            rows.append(
                {
                    "source_file": path.name,
                    "order_id": str(row[order_idx] or "").strip(),
                    "seller_liquidation_number": str(row[liquidation_idx] or "").strip(),
                    "payable_amount": parse_number(row[payable_idx]),
                }
            )
        workbook.close()
    return rows


def load_xml() -> list[XmlDoc]:
    docs: list[XmlDoc] = []
    for path in sorted(BASE_DIR.glob("XML/*.xml")):
        root = ET.parse(path).getroot()
        docs.append(
            XmlDoc(
                file_name=path.name,
                dte_type=root.findtext(".//sii:TipoDTE", namespaces=NS),
                folio=root.findtext(".//sii:Folio", namespaces=NS),
                issue_date=root.findtext(".//sii:FchEmis", namespaces=NS),
                total=parse_number(root.findtext(".//sii:MntTotal", namespaces=NS)),
                net=parse_number(root.findtext(".//sii:MntNeto", namespaces=NS)),
                vat=parse_number(root.findtext(".//sii:IVA", namespaces=NS)),
                commission_net=parse_number(root.findtext(".//sii:ValComNeto", namespaces=NS)),
                commission_vat=parse_number(root.findtext(".//sii:ValComIVA", namespaces=NS)),
                recipient_name=root.findtext(".//sii:RznSocRecep", namespaces=NS),
                detail_count=len(root.findall(".//sii:Detalle", NS)),
            )
        )
    return docs


def build_report() -> dict[str, object]:
    th = load_th()
    ciclos = load_ciclos()
    ff = load_ff()
    ff_adj = load_ff_adjustments()
    seller = load_seller()
    xml_docs = load_xml()

    th_orders = {row["order_id"] for row in th if row["order_id"]}
    th_invoices = {row["invoice_number"] for row in th if row["invoice_number"]}
    ciclo_orders = {row["order_id"] for row in ciclos if row["order_id"]}
    ciclo_invoices = {row["invoice_number"] for row in ciclos if row["invoice_number"]}
    ff_orders = {row["order_id"] for row in ff if row["order_id"]}
    ff_adj_orders = {row["order_id"] for row in ff_adj if row["order_id"]}
    seller_orders = {row["order_id"] for row in seller if row["order_id"]}
    seller_liquidations = {row["seller_liquidation_number"] for row in seller if row["seller_liquidation_number"]}
    xml_folios = {doc.folio for doc in xml_docs if doc.folio}

    cycle_by_invoice = defaultdict(set)
    for row in ciclos:
        if row["invoice_number"]:
            cycle_by_invoice[row["invoice_number"]].add(row["order_id"])

    th_by_invoice = defaultdict(set)
    for row in th:
        if row["invoice_number"]:
            th_by_invoice[row["invoice_number"]].add(row["order_id"])

    xml_type_counter = Counter(doc.dte_type for doc in xml_docs)
    xml_totals_by_type = defaultdict(float)
    for doc in xml_docs:
        xml_totals_by_type[doc.dte_type or "UNKNOWN"] += doc.total

    top_th_types = Counter(row["transaction_type"] for row in th).most_common(12)

    traceability = {
        "th_to_ciclos_by_invoice_matches": len(th_invoices & ciclo_invoices),
        "th_to_ciclos_by_order_matches": len(th_orders & ciclo_orders),
        "ciclos_to_seller_by_order_matches": len(ciclo_orders & seller_orders),
        "ff_to_seller_by_order_matches": len(ff_orders & seller_orders),
        "ff_adjustments_to_seller_by_order_matches": len(ff_adj_orders & seller_orders),
        "seller_to_xml_by_document_matches": len(seller_liquidations & xml_folios),
    }

    gaps = [
        "No existe llave directa entre `SELLER.Número documento liquidación` y `XML.Folio` en los archivos disponibles.",
        "Los CSV de `TH` y `CICLOS` presentan encabezados con codificación inconsistente; la normalización es obligatoria antes de conciliar.",
        "La cobertura temporal no es uniforme: `TH` tiene 8 archivos mensuales, `SELLER` 48 libros, `CICLOS` 49 cierres y `XML` 407 DTE.",
    ]

    lineage = [
        {
            "step": 1,
            "source": "TH",
            "role": "Libro mayor granular por evento",
            "join_key": "order_id, invoice_number, order_line_id, cycle_date",
        },
        {
            "step": 2,
            "source": "CICLOS",
            "role": "Resumen por orden dentro del ciclo/factura",
            "join_key": "invoice_number y order_id",
        },
        {
            "step": 3,
            "source": "SELLER",
            "role": "Liquidación comercial consolidada a pagar",
            "join_key": "order_id; document number sólo interno en estos datos",
        },
        {
            "step": 4,
            "source": "FF",
            "role": "Órdenes fulfillment y descuentos operacionales",
            "join_key": "order_id y order_line_id",
        },
        {
            "step": 5,
            "source": "XML",
            "role": "Certificación tributaria SII (DTE 33, 43, 52, 61)",
            "join_key": "folio fiscal; no aparece un puente directo a SELLER",
        },
    ]

    return {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "inventory": {
            "TH_files": len(list(BASE_DIR.glob("TH/*.csv"))),
            "CICLOS_files": len(list(BASE_DIR.glob("CICLOS/*.csv"))),
            "FF_files": len(list(BASE_DIR.glob("FF/*.csv"))),
            "SELLER_files": len(list(BASE_DIR.glob("SELLER/*.xlsx"))),
            "XML_files": len(list(BASE_DIR.glob("XML/*.xml"))),
        },
        "row_counts": {
            "TH_rows": len(th),
            "CICLOS_rows": len(ciclos),
            "FF_rows": len(ff),
            "FF_adjustment_rows": len(ff_adj),
            "SELLER_rows": len(seller),
        },
        "unique_keys": {
            "TH_orders": len(th_orders),
            "TH_invoices": len(th_invoices),
            "CICLOS_orders": len(ciclo_orders),
            "CICLOS_invoices": len(ciclo_invoices),
            "FF_orders": len(ff_orders),
            "SELLER_orders": len(seller_orders),
            "SELLER_liquidation_numbers": len(seller_liquidations),
            "XML_folios": len(xml_folios),
        },
        "traceability": traceability,
        "xml_summary": {
            "document_count_by_type": dict(xml_type_counter),
            "document_total_by_type": {key: round(value, 2) for key, value in xml_totals_by_type.items()},
        },
        "top_th_transaction_types": top_th_types,
        "lineage": lineage,
        "gaps": gaps,
        "sample_invoice_overlap": [
            {
                "invoice_number": invoice,
                "th_orders": sorted(th_by_invoice[invoice])[:5],
                "ciclos_orders": sorted(cycle_by_invoice[invoice])[:5],
            }
            for invoice in sorted(th_invoices & ciclo_invoices)[:10]
        ],
    }


def write_markdown(report: dict[str, object]) -> None:
    lines = [
        "# RIPLEY Traceability Report",
        "",
        f"- Generated at: {report['generated_at']}",
        f"- Files: {json.dumps(report['inventory'], ensure_ascii=False)}",
        f"- Rows: {json.dumps(report['row_counts'], ensure_ascii=False)}",
        "",
        "## End-to-End Lineage",
    ]
    for step in report["lineage"]:
        lines.append(
            f"- {step['step']}. {step['source']}: {step['role']} | Join: {step['join_key']}"
        )
    lines.extend(
        [
            "",
            "## Traceability Coverage",
            f"- {json.dumps(report['traceability'], ensure_ascii=False)}",
            "",
            "## XML Tax Documents",
            f"- {json.dumps(report['xml_summary'], ensure_ascii=False)}",
            "",
            "## Main Gaps",
        ]
    )
    for gap in report["gaps"]:
        lines.append(f"- {gap}")
    lines.extend(
        [
            "",
            "## TH Top Transaction Types",
        ]
    )
    for tx_type, count in report["top_th_transaction_types"]:
        lines.append(f"- {tx_type}: {count}")
    (BASE_DIR / "traceability_report.md").write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    report = build_report()
    (BASE_DIR / "traceability_report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    write_markdown(report)
    print(json.dumps(report["traceability"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
